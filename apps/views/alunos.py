from .utilitarios import *

import logging

from django.db import transaction
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST

from apps.forms import (
    AlunoEditForm,
    AlunoForm,
    PEDAGOGA_NAO_DEFINIDA,
    TIPOS_OCORRENCIA_PERMITIDOS,
    normalizar_tipo_ocorrencia,
    normalizar_turma_mapeamento,
    pedagoga_por_turma,
    resolver_pedagoga_turma,
)

logger = logging.getLogger(__name__)

MENSAGEM_OCORRENCIA_RELATORIO_SEMANAL = (
    '✅ Ocorrência registrada com sucesso. O responsável será incluído no relatório semanal, '
    'caso o aluno atinja o critério de envio.'
)
TIPOS_OCORRENCIA_WHATSAPP = {'atraso', 'piercing', 'desvio_normas', 'cabelo', 'uniforme', 'fora_sala', 'matando_aula'}
TIPOS_OCORRENCIA_REINCIDENCIA = {'piercing', 'desvio_normas', 'cabelo', 'uniforme'}
TIPOS_OCORRENCIA_LABELS = {
    'atraso': 'Atraso',
    'piercing': 'Uso de Piercing',
    'desvio_normas': 'Desvio de Normas',
    'cabelo': 'Cabelo',
    'uniforme': 'Uniforme',
    'fora_sala': 'Fora da sala',
    'matando_aula': 'Matando aula',
}


def _telefone_whatsapp_ocorrencia(aluno):
    for campo in ('telefone_responsavel', 'telefone'):
        telefone = getattr(aluno, campo, None)
        if telefone is None:
            continue
        telefone = str(telefone).strip()
        if telefone:
            return telefone
    return ''


def _data_ocorrencia_whatsapp(data_ocorrencia):
    try:
        return datetime.strptime(str(data_ocorrencia), '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return str(data_ocorrencia)


def _enviar_aviso_ocorrencia_whatsapp_digitador(request, aluno, tipo_ocorrencia, data_ocorrencia, mensagem_sucesso=None):
    messages.success(request, mensagem_sucesso or MENSAGEM_OCORRENCIA_RELATORIO_SEMANAL)


@login_required
@user_passes_test(grupo_digitadores, login_url='/')
def formulario_digitador(request):
    """View exclusiva para digitadores (apenas o formulario de ocorrencias)."""
    tipos_ocorrencia_permitidos = TIPOS_OCORRENCIA_PERMITIDOS

    # Recupera ultima data da sessao
    ultima_data_str = request.session.get('ultima_data_ocorrencia')
    if ultima_data_str:
        try:
            ultima_data = datetime.strptime(ultima_data_str, '%Y-%m-%d').date()
        except ValueError:
            ultima_data = timezone.now().date()
    else:
        ultima_data = timezone.now().date()

    # Recupera ultima turma da sessao
    ultima_turma_id = request.session.get('ultima_turma_ocorrencia_id')
    ultima_turma = None
    if ultima_turma_id:
        try:
            ultima_turma = Turma.objects.get(id=ultima_turma_id)
        except Turma.DoesNotExist:
            ultima_turma = None

    if request.method == 'POST':
        tipo_ocorrencia_postado = normalizar_tipo_ocorrencia(request.POST.get('tipo_ocorrencia'))
        if tipo_ocorrencia_postado not in tipos_ocorrencia_permitidos:
            messages.error(request, 'Tipo de ocorrencia invalido. Faltas devem ser registradas em Faltas de Alunos.')
            return redirect('formulario_digitador')

        form = RegistroOcorrenciaForm(request.POST)
        if form.is_valid():
            turma = form.cleaned_data['turma']
            numero = form.cleaned_data['numero_aluno']

            aluno = Aluno.objects.filter(turma=turma, numero=numero).first()
            if not aluno:
                messages.error(request, f'Aluno numero {numero} nao encontrado na turma {turma.nome}!')
                return render(request, 'ocorrencias/formulario_digitador.html', {
                    'form': form,
                    'ultimas_ocorrencias': RegistroOcorrenciaAluno.objects.select_related('aluno', 'aluno__turma').order_by('-data', '-horario_chegada')[:10]
                })

            ocorrencia = form.save(commit=False)

            # Pega o tipo de ocorrencia do formulario
            tipo_ocorrencia = tipo_ocorrencia_postado
            if tipo_ocorrencia not in tipos_ocorrencia_permitidos:
                messages.error(request, 'Tipo de ocorrencia invalido. Faltas devem ser registradas em Faltas de Alunos.')
                return redirect('formulario_digitador')
            ocorrencia.tipo_ocorrencia = tipo_ocorrencia
            ocorrencia.faltou = False

            # Pega o turno do formulario
            ocorrencia.turno = form.cleaned_data.get('turno', 'manha')
            pedagoga_informada = request.POST.get('atendido_por', '')
            logger.info('Turma normalizada para mapeamento de pedagoga: %s', normalizar_turma_mapeamento(aluno.turma))
            ocorrencia.atendido_por = resolver_pedagoga_turma(
                aluno.turma,
                pedagoga_informada,
            )

            if ocorrencia.horario_chegada == '':
                ocorrencia.horario_chegada = None
            if ocorrencia.horario_contato == '':
                ocorrencia.horario_contato = None

            ocorrencia.aluno = aluno
            ocorrencia.registrado_por = request.user
            try:
                with transaction.atomic():
                    ocorrencia.save()
            except IntegrityError:
                messages.error(request, 'Ja existe um registro para este aluno nesta data com o mesmo tipo de ocorrencia. Nao e possivel duplicar sem alterar a restricao atual do banco.')
                return redirect('formulario_digitador')

            request.session['ultima_data_ocorrencia'] = ocorrencia.data.isoformat()
            request.session['ultima_turma_ocorrencia_id'] = turma.id
            request.session['ultima_turma_ocorrencia_nome'] = turma.nome

            messages.success(request, f'{MENSAGEM_OCORRENCIA_RELATORIO_SEMANAL} Aluno: {aluno.nome} - Turma: {turma.nome}, No {numero}.')
            return redirect('formulario_digitador')
        else:
            messages.error(request, 'Erro no formulario. Verifique os dados.')
    else:
        initial_data = {
            'data': ultima_data.isoformat(),
            'faltou': False,
        }
        form = RegistroOcorrenciaForm(initial=initial_data)
        if ultima_turma:
            form.fields['turma'].initial = ultima_turma

    ultimas_ocorrencias = RegistroOcorrenciaAluno.objects.select_related('aluno', 'aluno__turma').order_by('-data', '-horario_chegada')[:10]

    return render(request, 'ocorrencias/formulario_digitador.html', {
        'form': form,
        'ultimas_ocorrencias': ultimas_ocorrencias
    })

# =============================================================================
# IMPORTAÃ‡ÃƒO DE ALUNOS VIA EXCEL
# =============================================================================
from openpyxl import load_workbook
from ..models import Turma, Aluno, AtestadoAluno

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def importar_alunos_excel(request):
    if request.method == 'POST' and request.FILES.get('arquivo'):
        arquivo = request.FILES['arquivo']
        wb = load_workbook(arquivo)
        ws = wb.active

        contador = 0
        erros = 0

        for row in ws.iter_rows(min_row=2, values_only=True):  # Pula cabeÃ§alho
            try:
                turma_nome = str(row[0]).strip() if row[0] else ''  # Coluna A: Turma (ex: 3A)
                numero = int(row[1]) if row[1] else 0               # Coluna B: NÃºmero
                nome = str(row[2]).strip() if row[2] else ''        # Coluna C: Nome do aluno

                if turma_nome and numero and nome:
                    # Extrai a sÃ©rie do nome da turma (ex: 3A -> 3Âº Ano)
                    serie = f"{turma_nome[0]}Âº Ano"

                    # Busca ou cria a turma
                    turma, created = Turma.objects.get_or_create(
                        nome=turma_nome,
                        defaults={
                            'serie': serie,
                            'ano': 2026,
                            'ativa': True
                        }
                    )

                    # Cria o aluno
                    Aluno.objects.create(
                        nome=nome,
                        numero=numero,
                        turma=turma
                    )
                    contador += 1
                else:
                    erros += 1
            except Exception as e:
                erros += 1
                print(f"Erro na linha: {e}")

        messages.success(request, f'{contador} alunos importados com sucesso!')
        if erros > 0:
            messages.warning(request, f'{erros} linhas com erro foram ignoradas.')

        return redirect('controle_faltas_alunos')

    return render(request, 'importar_alunos.html')


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def cadastrar_aluno(request):
    if request.method == 'POST':
        form = AlunoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'âœ… Aluno cadastrado com sucesso!')
            return redirect('painel_equipe')
    else:
        form = AlunoForm()

    return render(request, 'cadastrar_aluno.html', {'form': form})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def listar_alunos(request):
    alunos = Aluno.objects.all().order_by('turma', 'numero')
    return render(request, 'listar_alunos.html', {'alunos': alunos})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def atestados_alunos(request):
    if request.method == 'POST':
        turma_id = request.POST.get('turma')
        numero_aluno = request.POST.get('numero_aluno', '').strip()
        data_inicio = parse_date(request.POST.get('data_inicio', ''))
        data_fim = parse_date(request.POST.get('data_fim', ''))
        observacao = request.POST.get('observacao', '').strip()
        ativo = request.POST.get('ativo') == 'on'

        aluno = Aluno.objects.filter(turma_id=turma_id, numero=numero_aluno, ativo=True).select_related('turma').first()
        if not aluno:
            messages.error(request, 'Aluno ativo nao encontrado para a turma e numero informados.')
            return redirect('atestados_alunos')

        if not data_inicio or not data_fim:
            messages.error(request, 'Informe data inicio e data fim.')
            return redirect('atestados_alunos')

        if data_fim < data_inicio:
            messages.error(request, 'Data fim nao pode ser menor que data inicio.')
            return redirect('atestados_alunos')

        AtestadoAluno.objects.create(
            aluno=aluno,
            data_inicio=data_inicio,
            data_fim=data_fim,
            observacao=observacao,
            ativo=ativo,
            criado_por=request.user,
        )
        messages.success(request, f'Atestado cadastrado para {aluno.nome}.')
        return redirect('atestados_alunos')

    turmas = Turma.objects.filter(ativa=True).order_by('turno', 'nome')
    atestados = AtestadoAluno.objects.select_related(
        'aluno',
        'aluno__turma',
        'criado_por',
    ).order_by('-data_inicio', 'aluno__turma__nome', 'aluno__numero')

    return render(request, 'atestados_alunos.html', {
        'turmas': [
            {
                'id': turma.id,
                'nome': turma.nome,
                'serie': turma.serie,
                'turno': turma.turno or '',
            }
            for turma in turmas
        ],
        'atestados': atestados,
    })


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def buscar_aluno_atestado(request):
    turma_id = request.GET.get('turma_id')
    numero = request.GET.get('numero', '').strip()

    if not turma_id or not numero:
        return JsonResponse({'encontrado': False, 'erro': 'Informe turma e numero.'})

    aluno = Aluno.objects.filter(turma_id=turma_id, numero=numero, ativo=True).select_related('turma').first()
    if not aluno:
        return JsonResponse({'encontrado': False, 'erro': 'Aluno ativo nao encontrado nesta turma.'})

    return JsonResponse({
        'encontrado': True,
        'nome': aluno.nome,
        'numero': aluno.numero,
        'turma': aluno.turma.nome,
        'turma_label': str(aluno.turma),
        'turno': aluno.turma.turno or '',
        'turno_label': aluno.turma.get_turno_display() if aluno.turma.turno else '-',
    })


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def editar_atestado_aluno(request, atestado_id):
    atestado = get_object_or_404(AtestadoAluno.objects.select_related('aluno', 'aluno__turma'), id=atestado_id)

    if request.method == 'POST':
        data_inicio = parse_date(request.POST.get('data_inicio', ''))
        data_fim = parse_date(request.POST.get('data_fim', ''))

        if not data_inicio or not data_fim:
            messages.error(request, 'Informe data inicio e data fim.')
            return redirect('editar_atestado_aluno', atestado_id=atestado.id)

        if data_fim < data_inicio:
            messages.error(request, 'Data fim nao pode ser menor que data inicio.')
            return redirect('editar_atestado_aluno', atestado_id=atestado.id)

        atestado.data_inicio = data_inicio
        atestado.data_fim = data_fim
        atestado.observacao = request.POST.get('observacao', '').strip()
        atestado.ativo = request.POST.get('ativo') == 'on'
        atestado.save(update_fields=['data_inicio', 'data_fim', 'observacao', 'ativo'])
        messages.success(request, f'Atestado de {atestado.aluno.nome} atualizado.')
        return redirect('atestados_alunos')

    return render(request, 'editar_atestado_aluno.html', {'atestado': atestado})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
@require_POST
def alternar_atestado_aluno(request, atestado_id):
    atestado = get_object_or_404(AtestadoAluno, id=atestado_id)
    atestado.ativo = not atestado.ativo
    atestado.save(update_fields=['ativo'])
    status = 'ativado' if atestado.ativo else 'desativado'
    messages.success(request, f'Atestado {status} com sucesso.')
    return redirect('atestados_alunos')


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
@require_POST
def excluir_atestado_aluno(request, atestado_id):
    atestado = get_object_or_404(AtestadoAluno.objects.select_related('aluno'), id=atestado_id)
    if atestado.ativo:
        messages.error(request, 'Nao e possivel excluir um atestado ativo. Desative o atestado antes de excluir.')
        return redirect('atestados_alunos')

    aluno_nome = atestado.aluno.nome
    atestado.delete()
    messages.success(request, f'Atestado de {aluno_nome} excluido com sucesso.')
    return redirect('atestados_alunos')


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def editar_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)
    if request.method == 'POST':
        form = AlunoEditForm(request.POST, instance=aluno)
        if form.is_valid():
            form.save()
            messages.success(request, f'âœ… Aluno {aluno.nome} atualizado!')
            return redirect('listar_alunos')
    else:
        form = AlunoEditForm(instance=aluno)

    return render(request, 'editar_aluno.html', {'form': form, 'aluno': aluno})


# =============================================================================
# CADASTRO E EDIÃ‡ÃƒO DE ALUNOS (jÃ¡ existente)
# =============================================================================

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def editar_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)
    if request.method == 'POST':
        form = AlunoEditForm(request.POST, instance=aluno)
        if form.is_valid():
            form.save()
            messages.success(request, f'âœ… Aluno {aluno.nome} atualizado!')
            return redirect('listar_alunos')
    else:
        form = AlunoEditForm(instance=aluno)

    return render(request, 'editar_aluno.html', {'form': form, 'aluno': aluno})
